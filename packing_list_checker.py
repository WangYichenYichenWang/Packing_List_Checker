"""
Packing List Checker
====================
Combines supplier packing lists (Excel), checks them line by line and against
the commercial invoice, and writes one Excel report with four sheets:

  Summary        issue counts per check
  Issues         every problem found, one row per issue
  HS_Summary     cartons, quantity, weights and value per shipment, HS code and origin
  Packing_Lists  all packing-list lines combined, with the source file name

Usage:
    python packing_list_checker.py
    python packing_list_checker.py --packing-lists data/packing_lists \
        --invoice data/commercial_invoice.xlsx --out output/packing_list_report.xlsx
"""
import argparse
import re
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill

# Suppliers label the same column differently. Headers are first cleaned
# (lower case, no dots or brackets, spaces -> "_") and then mapped here.
COLUMN_ALIASES = {
    "shipment_no": "shipment_id", "shipment": "shipment_id",
    "item_no": "sku", "article_no": "sku",
    "goods_description": "description",
    "hs": "hs_code", "commodity_code": "hs_code", "tariff_code": "hs_code",
    "ctns": "cartons", "no_of_cartons": "cartons",
    "qty_per_ctn": "qty_per_carton", "pcs_per_ctn": "qty_per_carton",
    "qty": "total_qty", "quantity": "total_qty", "total_pcs": "total_qty",
    "nw_kg": "net_weight_kg", "net_weight": "net_weight_kg",
    "gw_kg": "gross_weight_kg", "gross_weight": "gross_weight_kg",
    "unit_price": "unit_price_eur",
    "origin": "origin_country", "country_of_origin": "origin_country",
}

PL_COLUMNS = ["shipment_id", "sku", "description", "hs_code", "cartons",
              "qty_per_carton", "total_qty", "net_weight_kg", "gross_weight_kg"]
INVOICE_COLUMNS = ["shipment_id", "sku", "hs_code", "total_qty", "unit_price_eur", "origin_country"]
NUMERIC_COLUMNS = ["cartons", "qty_per_carton", "total_qty", "net_weight_kg",
                   "gross_weight_kg", "unit_price_eur"]
VALID_HS_LENGTHS = (6, 8, 10)  # HS (6), EU Combined Nomenclature (8), TARIC (10)


# ---------------------------------------------------------------- loading ---

def clean_header(name):
    name = re.sub(r"[().]", "", str(name).strip().lower())
    name = re.sub(r"\s+", "_", name)
    return COLUMN_ALIASES.get(name, name)


def clean_hs_code(value):
    """'9403.50' -> '940350'; empty cells -> ''."""
    if pd.isna(value):
        return ""
    return re.sub(r"\D", "", str(value))


def load_table(path, columns):
    """Read one Excel file and return it with standard column names and types."""
    df = pd.read_excel(path, dtype=str).dropna(how="all")
    df = df.rename(columns=clean_header)

    missing = [c for c in columns if c not in df.columns]
    if missing:
        raise ValueError(f"{path.name}: columns not found {missing}. "
                         f"Add the supplier's header to COLUMN_ALIASES.")

    df = df[columns].copy()
    df["shipment_id"] = df["shipment_id"].str.strip()
    df["sku"] = df["sku"].str.strip()
    df["hs_code"] = df["hs_code"].map(clean_hs_code)
    for col in NUMERIC_COLUMNS:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")
    df["source_file"] = path.name
    return df


def load_packing_lists(folder):
    files = sorted(Path(folder).glob("*.xlsx"))
    if not files:
        raise FileNotFoundError(f"No .xlsx packing lists found in {folder}")
    return pd.concat([load_table(f, PL_COLUMNS) for f in files], ignore_index=True)


# ----------------------------------------------------------------- checks ---

def make_issue(severity, row, check, detail):
    return {"severity": severity, "shipment_id": row.shipment_id, "sku": row.sku,
            "check": check, "detail": detail, "source_file": row.source_file}


def check_packing_lists(pl):
    """Checks that only need the packing list itself."""
    issues = []
    for row in pl.itertuples(index=False):
        # 1. Required values present
        for col in ["hs_code", "cartons", "qty_per_carton", "total_qty",
                    "net_weight_kg", "gross_weight_kg"]:
            value = getattr(row, col)
            if value == "" or pd.isna(value):
                issues.append(make_issue("Error", row, "Missing value", f"'{col}' is empty"))

        # 2. HS code has a valid length
        if row.hs_code and len(row.hs_code) not in VALID_HS_LENGTHS:
            issues.append(make_issue("Error", row, "Invalid HS code",
                                     f"'{row.hs_code}' has {len(row.hs_code)} digits "
                                     f"(expected 6, 8 or 10)"))

        # 3. Cartons x quantity per carton = total quantity
        if pd.notna(row.cartons) and pd.notna(row.qty_per_carton) and pd.notna(row.total_qty):
            expected = row.cartons * row.qty_per_carton
            if expected != row.total_qty:
                issues.append(make_issue("Error", row, "Quantity arithmetic",
                                         f"{row.cartons:g} cartons x {row.qty_per_carton:g} pcs "
                                         f"= {expected:g}, but total qty is {row.total_qty:g}"))

        # 4. Gross weight must not be below net weight
        if pd.notna(row.net_weight_kg) and pd.notna(row.gross_weight_kg):
            if row.gross_weight_kg < row.net_weight_kg:
                issues.append(make_issue("Error", row, "Weight logic",
                                         f"Gross weight {row.gross_weight_kg:g} kg is below "
                                         f"net weight {row.net_weight_kg:g} kg"))

    # 5. Same item listed twice in one shipment (can be legitimate, so a warning)
    duplicates = pl[pl.duplicated(subset=["shipment_id", "sku"], keep="first")]
    for row in duplicates.itertuples(index=False):
        issues.append(make_issue("Warning", row, "Possible duplicate line",
                                 "Item appears more than once in this shipment"))
    return issues


def check_against_invoice(pl, invoice, invoice_name):
    """Compare packing-list totals with the commercial invoice per shipment and item."""
    pl_totals = pl.groupby(["shipment_id", "sku"], as_index=False).agg(
        pl_qty=("total_qty", "sum"), pl_hs=("hs_code", "first"), source_file=("source_file", "first"))
    inv_totals = invoice.groupby(["shipment_id", "sku"], as_index=False).agg(
        inv_qty=("total_qty", "sum"), inv_hs=("hs_code", "first"))
    merged = pl_totals.merge(inv_totals, on=["shipment_id", "sku"], how="outer", indicator="match")

    issues = []
    for row in merged.itertuples(index=False):
        if row.match == "left_only":
            issues.append(make_issue("Error", row, "Not on invoice",
                                     "Packed but missing from the commercial invoice "
                                     "(must still be declared, even if free of charge)"))
        elif row.match == "right_only":
            row = row._replace(source_file=invoice_name)
            issues.append(make_issue("Error", row, "Not in packing list",
                                     "Invoiced but not found in any packing list"))
        else:
            if row.pl_qty != row.inv_qty:
                issues.append(make_issue("Error", row, "Quantity mismatch",
                                         f"Packing list {row.pl_qty:g} pcs vs invoice {row.inv_qty:g} pcs"))
            if row.pl_hs and row.inv_hs and row.pl_hs != row.inv_hs:
                issues.append(make_issue("Error", row, "HS code mismatch",
                                         f"Packing list {row.pl_hs} vs invoice {row.inv_hs}"))
    return issues


# ---------------------------------------------------------------- outputs ---

def build_hs_summary(pl, invoice):
    """Totals per shipment, HS code and origin: the figures a customs declaration needs.
    Value = packing-list quantity x invoice unit price."""
    prices = invoice.groupby(["shipment_id", "sku"], as_index=False).agg(
        unit_price_eur=("unit_price_eur", "first"), origin_country=("origin_country", "first"))
    lines = pl.merge(prices, on=["shipment_id", "sku"], how="left")
    lines["value_eur"] = lines["total_qty"] * lines["unit_price_eur"]
    lines["hs_code"] = lines["hs_code"].replace("", "MISSING")
    lines["origin_country"] = lines["origin_country"].fillna("NOT ON INVOICE")

    summary = lines.groupby(["shipment_id", "hs_code", "origin_country"], as_index=False).agg(
        lines=("sku", "count"),
        cartons=("cartons", "sum"),
        total_qty=("total_qty", "sum"),
        net_weight_kg=("net_weight_kg", "sum"),
        gross_weight_kg=("gross_weight_kg", "sum"),
        value_eur=("value_eur", lambda s: s.sum(min_count=1)),
    )
    return summary.round(2)


def build_summary(pl, issues):
    overview = pd.DataFrame({
        "item": ["Packing list files", "Shipments", "Packing list lines", "Errors", "Warnings"],
        "count": [pl["source_file"].nunique(), pl["shipment_id"].nunique(), len(pl),
                  int((issues["severity"] == "Error").sum()) if len(issues) else 0,
                  int((issues["severity"] == "Warning").sum()) if len(issues) else 0],
    })
    if len(issues):
        by_check = issues.groupby("check").size().reset_index(name="count").rename(columns={"check": "item"})
        overview = pd.concat([overview, pd.DataFrame([{"item": "", "count": None}]), by_check],
                             ignore_index=True)
    overview["count"] = overview["count"].astype("Int64")
    return overview


def format_workbook(writer):
    header_fill = PatternFill("solid", fgColor="1F3864")
    fills = {"Error": PatternFill("solid", fgColor="F8D7DA"),
             "Warning": PatternFill("solid", fgColor="FFF3CD")}
    for ws in writer.book.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for column in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in column)
            ws.column_dimensions[column[0].column_letter].width = min(width + 2, 70)
        if ws.title == "Issues":
            for row in ws.iter_rows(min_row=2):
                fill = fills.get(row[0].value)
                if fill:
                    for cell in row:
                        cell.fill = fill


def write_report(path, summary, issues, hs_summary, pl):
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        issues.to_excel(writer, sheet_name="Issues", index=False)
        hs_summary.to_excel(writer, sheet_name="HS_Summary", index=False)
        pl.to_excel(writer, sheet_name="Packing_Lists", index=False)
        format_workbook(writer)


# ------------------------------------------------------------------- main ---

def main():
    parser = argparse.ArgumentParser(description="Check packing lists against the commercial invoice.")
    parser.add_argument("--packing-lists", default="data/packing_lists", help="folder of .xlsx packing lists")
    parser.add_argument("--invoice", default="data/commercial_invoice.xlsx", help="commercial invoice .xlsx")
    parser.add_argument("--out", default="output/packing_list_report.xlsx", help="report to write")
    args = parser.parse_args()

    invoice_path = Path(args.invoice)
    pl = load_packing_lists(args.packing_lists)
    invoice = load_table(invoice_path, INVOICE_COLUMNS)

    issue_list = check_packing_lists(pl) + check_against_invoice(pl, invoice, invoice_path.name)
    issues = pd.DataFrame(issue_list, columns=["severity", "shipment_id", "sku", "check",
                                               "detail", "source_file"])
    issues = issues.sort_values(["severity", "shipment_id", "sku"], ignore_index=True)

    summary = build_summary(pl, issues)
    hs_summary = build_hs_summary(pl, invoice)
    write_report(Path(args.out), summary, issues, hs_summary, pl)

    print(f"Checked {len(pl)} lines from {pl['source_file'].nunique()} packing lists.")
    print(f"Found {(issues['severity'] == 'Error').sum()} errors and "
          f"{(issues['severity'] == 'Warning').sum()} warnings.")
    print(f"Report written to {args.out}")


if __name__ == "__main__":
    main()
